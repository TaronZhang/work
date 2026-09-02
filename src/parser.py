"""文档解析模块 — 支持 .doc / .docx / .pdf / .txt"""

from __future__ import annotations

import os
import shutil
import subprocess
from dataclasses import dataclass, field
from typing import Optional
from src.utils import detect_encoding, normalize_text


@dataclass
class Paragraph:
    """段落数据结构"""
    text: str
    style: str = 'Normal'     # 段落样式: Normal, Heading1, table, etc.
    page: Optional[int] = None   # 所在页码
    index: int = 0            # 全局段落序号


@dataclass
class ParsedDocument:
    """解析后的文档"""
    filename: str
    full_text: str
    paragraphs: list[Paragraph] = field(default_factory=list)
    metadata: dict = field(default_factory=dict)


def parse_file(filepath: str, original_filename: str = None) -> ParsedDocument:
    """解析文件入口，根据扩展名分派"""
    ext = os.path.splitext(filepath)[1].lower()
    filename = original_filename or os.path.basename(filepath)

    if ext == '.docx':
        return parse_docx(filepath, filename)
    elif ext == '.doc':
        return parse_doc(filepath, filename)
    elif ext == '.pdf':
        return parse_pdf(filepath, filename)
    elif ext == '.txt':
        return parse_txt(filepath, filename)
    else:
        raise ValueError(f"不支持的文件格式: {ext}")


def parse_doc(filepath: str, filename: str) -> ParsedDocument:
    """解析 .doc 文件（旧版 Word 二进制格式）

    优先使用 antiword，回退到 catdoc，两者均不可用时尝试 olefile。
    """
    raw_text = _extract_doc_text(filepath)

    paragraphs = []
    idx = 0
    for line in raw_text.split('\n'):
        line = line.strip()
        if line:
            paragraphs.append(Paragraph(text=line, style='Normal', index=idx))
            idx += 1

    full_text = normalize_text(raw_text)

    return ParsedDocument(
        filename=filename,
        full_text=full_text,
        paragraphs=paragraphs,
        metadata={
            'format': 'doc',
            'chars': len(full_text),
            'paragraphs': len(paragraphs),
            'extractor': _detect_doc_extractor(),
        }
    )


def _detect_doc_extractor() -> str:
    """检测可用的 .doc 解析工具"""
    for tool in ('antiword', 'catdoc'):
        if shutil.which(tool):
            return tool
    return 'olefile'


def _extract_doc_text(filepath: str) -> str:
    """从 .doc 文件提取纯文本"""
    for tool, args in [
        ('antiword', ['-m', 'UTF-8.txt']),
        ('catdoc', ['-d', 'utf-8']),
    ]:
        if shutil.which(tool):
            try:
                result = subprocess.run(
                    [tool] + args + [filepath],
                    capture_output=True, text=True, timeout=60,
                )
                if result.returncode == 0 and result.stdout.strip():
                    return result.stdout
            except (subprocess.TimeoutExpired, FileNotFoundError):
                continue

    try:
        return _extract_doc_olefile(filepath)
    except Exception:
        raise RuntimeError(
            "无法解析 .doc 文件，请安装 antiword 或 catdoc："
            " apt-get install antiword catdoc"
        )


def _extract_doc_olefile(filepath: str) -> str:
    """使用 olefile 作为最后兜底提取 .doc 文本"""
    import olefile

    ole = olefile.OleFileIO(filepath)
    try:
        stream = ole.openstream('WordDocument')
        data = stream.read()
        text = data.decode('utf-16-le', errors='ignore')
        clean = ''.join(ch for ch in text if ch.isprintable() or ch in '\n\r\t')
        return clean
    finally:
        ole.close()


def parse_docx(filepath: str, filename: str) -> ParsedDocument:
    """解析 .docx 文件"""
    from docx import Document

    doc = Document(filepath)
    paragraphs = []
    idx = 0

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            style = para.style.name if para.style else 'Normal'
            paragraphs.append(Paragraph(text=text, style=style, index=idx))
            idx += 1

    # 也提取表格中的文字
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                row_text = ' | '.join(cells)
                paragraphs.append(Paragraph(text=row_text, style='table', index=idx))
                idx += 1

    full_text = '\n'.join(p.text for p in paragraphs)
    full_text = normalize_text(full_text)

    return ParsedDocument(
        filename=filename,
        full_text=full_text,
        paragraphs=paragraphs,
        metadata={
            'format': 'docx',
            'chars': len(full_text),
            'paragraphs': len(paragraphs),
            'tables': len(doc.tables),
        }
    )


def parse_pdf(filepath: str, filename: str) -> ParsedDocument:
    """解析 .pdf 文件"""
    import pdfplumber

    paragraphs = []
    idx = 0

    with pdfplumber.open(filepath) as pdf:
        for page_num, page in enumerate(pdf.pages, 1):
            text = page.extract_text()
            if text:
                for line in text.strip().split('\n'):
                    line = line.strip()
                    if line:
                        paragraphs.append(Paragraph(
                            text=line,
                            style='Normal',
                            page=page_num,
                            index=idx
                        ))
                        idx += 1

            # 也提取表格
            for table in page.extract_tables():
                for row in table:
                    cells = [c.strip() if c else '' for c in row]
                    row_text = ' | '.join(c for c in cells if c)
                    if row_text:
                        paragraphs.append(Paragraph(
                            text=row_text,
                            style='table',
                            page=page_num,
                            index=idx
                        ))
                        idx += 1

    full_text = '\n'.join(p.text for p in paragraphs)
    full_text = normalize_text(full_text)

    return ParsedDocument(
        filename=filename,
        full_text=full_text,
        paragraphs=paragraphs,
        metadata={
            'format': 'pdf',
            'chars': len(full_text),
            'paragraphs': len(paragraphs),
            'pages': len(pdf.pages) if hasattr(pdf, 'pages') else '?',
        }
    )


def parse_txt(filepath: str, filename: str) -> ParsedDocument:
    """解析 .txt 文件（自动检测编码）"""
    encoding = detect_encoding(filepath)

    with open(filepath, 'r', encoding=encoding, errors='replace') as f:
        raw_text = f.read()

    paragraphs = []
    idx = 0
    for line in raw_text.split('\n'):
        line = line.strip()
        if line:
            paragraphs.append(Paragraph(text=line, style='Normal', index=idx))
            idx += 1

    full_text = normalize_text(raw_text)

    return ParsedDocument(
        filename=filename,
        full_text=full_text,
        paragraphs=paragraphs,
        metadata={
            'format': 'txt',
            'encoding': encoding,
            'chars': len(full_text),
            'paragraphs': len(paragraphs),
        }
    )
