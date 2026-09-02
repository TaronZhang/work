"""Export module — generate DOCX and PDF reports from comparison results."""

import io
from datetime import datetime

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak,
)


def _severity_color(level: str) -> str:
    if level == 'high':
        return '#f04a4a'
    if level == 'medium':
        return '#f0c040'
    return '#3dd68c'


def _severity_color_rgb(level: str):
    if level == 'high':
        return RGBColor(0xF0, 0x4A, 0x4A)
    if level == 'medium':
        return RGBColor(0xF0, 0xC0, 0x40)
    return RGBColor(0x3D, 0xD6, 0x8C)


def _severity_color_rl(level: str):
    if level == 'high':
        return colors.HexColor('#f04a4a')
    if level == 'medium':
        return colors.HexColor('#f0c040')
    return colors.HexColor('#3dd68c')


def generate_docx(results: list, mode: str, filter_info: dict | None = None) -> bytes:
    """Generate a DOCX report from comparison results.

    When mode='ai' and results contain ai_analysis, the report includes:
    - Adjusted duplicate percentage alongside original
    - AI summary text
    - Per-segment classification (合规应答/实际重复) and judgment reason
    """
    doc = Document()

    style = doc.styles['Normal']
    style.font.name = 'SimSun'
    style.font.size = Pt(11)

    title = doc.add_heading('文档重复度比对报告', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f'生成时间：{datetime.now().strftime("%Y-%m-%d %H:%M:%S")}    模式：{"AI 辅助" if mode == "ai" else "快速"}')
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    if filter_info and filter_info.get('categories'):
        doc.add_heading('AI 过滤摘要', level=1)
        for cat in filter_info['categories']:
            doc.add_paragraph(
                f'{cat["label"]}：{cat["count"]} 段',
                style='List Bullet',
            )

    doc.add_heading('比对结果摘要', level=1)

    has_ai = mode == 'ai' and any(r.get('ai_analysis') for r in results)
    ncols = 6 if has_ai else 5
    table = doc.add_table(rows=1, cols=ncols)
    table.style = 'Light Grid Accent 1'
    hdr = table.rows[0].cells
    hdr[0].text = '文档 A'
    hdr[1].text = '文档 B'
    hdr[2].text = '原始重复率'
    if has_ai:
        hdr[3].text = '调整后重复率'
        hdr[4].text = '严重程度'
        hdr[5].text = '匹配段数'
    else:
        hdr[3].text = '严重程度'
        hdr[4].text = '匹配段数'

    for r in results:
        row = table.add_row().cells
        row[0].text = r.get('label_a', '')
        row[1].text = r.get('label_b', '')
        pct = r.get('duplicate_percentage', 0)
        row[2].text = f'{pct:.1f}%'
        if has_ai:
            ai = r.get('ai_analysis')
            adj = ai['adjusted_percentage'] if ai else pct
            row[3].text = f'{adj:.1f}%'
            sev = r.get('severity', {})
            row[4].text = sev.get('label', '')
            row[5].text = str(len(r.get('segments', [])))
        else:
            sev = r.get('severity', {})
            row[3].text = sev.get('label', '')
            row[4].text = str(len(r.get('segments', [])))

    for r in results:
        ai = r.get('ai_analysis')
        pct = r.get('duplicate_percentage', 0)
        if ai:
            adj = ai['adjusted_percentage']
            doc.add_heading(
                f'{r.get("label_a", "A")} ↔ {r.get("label_b", "B")} — 原始 {pct:.1f}% → 调整后 {adj:.1f}%',
                level=2,
            )
            doc.add_paragraph(ai.get('summary', ''), style='Intense Quote')
        else:
            doc.add_heading(
                f'{r.get("label_a", "A")} ↔ {r.get("label_b", "B")} — 重复率 {pct:.1f}%',
                level=2,
            )

        meta = doc.add_paragraph()
        meta_run = meta.add_run(
            f'匹配字符：{r.get("matched_chars", 0)} / {r.get("total_tender_chars", 0)}    '
            f'严重程度：{r.get("severity", {}).get("label", "")}'
        )
        meta_run.font.size = Pt(9)
        meta_run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

        segments = r.get('segments', [])
        if not segments:
            doc.add_paragraph('未发现匹配的重复段落。')
            continue

        seg_cls_map = {}
        if ai:
            for sc in ai.get('segment_classifications', []):
                seg_cls_map[sc['segment_id']] = sc

        for seg in segments:
            sim_pct = round(seg.get('similarity', 0) * 100)
            seg_id = seg.get('id', 0)
            sc = seg_cls_map.get(seg_id)
            cls_label = ''
            if sc:
                cls_label = ' [合规应答]' if sc['classification'] == 'compliance_response' else ' [实际重复]'
            doc.add_heading(f'匹配段 #{seg_id + 1} — 相似度 {sim_pct}%{cls_label}', level=3)

            if sc and sc.get('reason'):
                reason_p = doc.add_paragraph()
                reason_run = reason_p.add_run(f'AI 判定理由：{sc["reason"]}')
                reason_run.font.size = Pt(9)
                reason_run.font.color.rgb = RGBColor(0x66, 0x66, 0x99)

            seg_table = doc.add_table(rows=2, cols=2)
            seg_table.style = 'Light List Accent 1'
            seg_table.rows[0].cells[0].text = r.get('label_a', '文档 A')
            seg_table.rows[0].cells[1].text = r.get('label_b', '文档 B')
            seg_table.rows[1].cells[0].text = seg.get('tender_text', '')
            seg_table.rows[1].cells[1].text = seg.get('bid_text', '')

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


def generate_pdf(results: list, mode: str, filter_info: dict | None = None) -> bytes:
    """Generate a PDF report from comparison results.

    When mode='ai' and results contain ai_analysis, the report includes:
    - Adjusted duplicate percentage alongside original
    - AI summary text
    - Per-segment classification (合规应答/实际重复) and judgment reason
    """
    pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        topMargin=20 * mm, bottomMargin=20 * mm,
        leftMargin=15 * mm, rightMargin=15 * mm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'CnTitle', parent=styles['Title'],
        fontName='STSong-Light', fontSize=20, leading=28,
        alignment=1,
    )
    heading_style = ParagraphStyle(
        'CnHeading', parent=styles['Heading1'],
        fontName='STSong-Light', fontSize=14, leading=20,
        spaceBefore=12, spaceAfter=6,
    )
    sub_heading_style = ParagraphStyle(
        'CnSubHeading', parent=styles['Heading2'],
        fontName='STSong-Light', fontSize=12, leading=16,
        spaceBefore=8, spaceAfter=4,
    )
    body_style = ParagraphStyle(
        'CnBody', parent=styles['Normal'],
        fontName='STSong-Light', fontSize=10, leading=14,
    )
    small_style = ParagraphStyle(
        'CnSmall', parent=styles['Normal'],
        fontName='STSong-Light', fontSize=8, leading=12,
        textColor=colors.grey,
    )
    cell_style = ParagraphStyle(
        'CnCell', parent=styles['Normal'],
        fontName='STSong-Light', fontSize=9, leading=12,
    )
    reason_style = ParagraphStyle(
        'CnReason', parent=styles['Normal'],
        fontName='STSong-Light', fontSize=9, leading=12,
        textColor=colors.HexColor('#5555aa'),
        leftIndent=10,
    )

    elems = []

    elems.append(Paragraph('文档重复度比对报告', title_style))
    elems.append(Spacer(1, 4 * mm))
    elems.append(Paragraph(
        f'生成时间：{datetime.now().strftime("%Y-%m-%d %H:%M:%S")}    '
        f'模式：{"AI 辅助" if mode == "ai" else "快速"}',
        small_style,
    ))
    elems.append(Spacer(1, 8 * mm))

    if filter_info and filter_info.get('categories'):
        elems.append(Paragraph('AI 过滤摘要', heading_style))
        for cat in filter_info['categories']:
            elems.append(Paragraph(
                f'{cat["label"]}：{cat["count"]} 段',
                body_style,
            ))
        elems.append(Spacer(1, 6 * mm))

    has_ai = mode == 'ai' and any(r.get('ai_analysis') for r in results)

    elems.append(Paragraph('比对结果摘要', heading_style))

    if has_ai:
        summary_data = [['文档 A', '文档 B', '原始重复率', '调整后重复率', '严重程度', '匹配段数']]
        for r in results:
            ai = r.get('ai_analysis')
            adj = ai['adjusted_percentage'] if ai else r.get('duplicate_percentage', 0)
            summary_data.append([
                Paragraph(r.get('label_a', ''), cell_style),
                Paragraph(r.get('label_b', ''), cell_style),
                f'{r.get("duplicate_percentage", 0):.1f}%',
                f'{adj:.1f}%',
                r.get('severity', {}).get('label', ''),
                str(len(r.get('segments', []))),
            ])
        col_widths = [35 * mm, 35 * mm, 18 * mm, 18 * mm, 22 * mm, 18 * mm]
    else:
        summary_data = [['文档 A', '文档 B', '重复率', '严重程度', '匹配段数']]
        for r in results:
            summary_data.append([
                Paragraph(r.get('label_a', ''), cell_style),
                Paragraph(r.get('label_b', ''), cell_style),
                f'{r.get("duplicate_percentage", 0):.1f}%',
                r.get('severity', {}).get('label', ''),
                str(len(r.get('segments', []))),
            ])
        col_widths = [45 * mm, 45 * mm, 20 * mm, 25 * mm, 20 * mm]

    summary_table = Table(summary_data, colWidths=col_widths, repeatRows=1)
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4a9eff')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'STSong-Light'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('ALIGN', (2, 0), (-1, -1), 'CENTER'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#444444')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f5f5f5')]),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    elems.append(summary_table)
    elems.append(Spacer(1, 10 * mm))

    for r in results:
        ai = r.get('ai_analysis')
        pct = r.get('duplicate_percentage', 0)
        if ai:
            adj = ai['adjusted_percentage']
            elems.append(Paragraph(
                f'{r.get("label_a", "A")} ↔ {r.get("label_b", "B")} — 原始 {pct:.1f}% → 调整后 {adj:.1f}%',
                heading_style,
            ))
            elems.append(Paragraph(
                escape_pdf_text(ai.get('summary', '')),
                body_style,
            ))
        else:
            elems.append(Paragraph(
                f'{r.get("label_a", "A")} ↔ {r.get("label_b", "B")} — 重复率 {pct:.1f}%',
                heading_style,
            ))
        elems.append(Paragraph(
            f'匹配字符：{r.get("matched_chars", 0)} / {r.get("total_tender_chars", 0)}    '
            f'严重程度：{r.get("severity", {}).get("label", "")}',
            small_style,
        ))
        elems.append(Spacer(1, 4 * mm))

        segments = r.get('segments', [])
        if not segments:
            elems.append(Paragraph('未发现匹配的重复段落。', body_style))
            elems.append(Spacer(1, 6 * mm))
            continue

        seg_cls_map = {}
        if ai:
            for sc in ai.get('segment_classifications', []):
                seg_cls_map[sc['segment_id']] = sc

        for seg in segments:
            sim_pct = round(seg.get('similarity', 0) * 100)
            seg_id = seg.get('id', 0)
            sc = seg_cls_map.get(seg_id)
            cls_label = ''
            if sc:
                cls_label = ' [合规应答]' if sc['classification'] == 'compliance_response' else ' [实际重复]'
            elems.append(Paragraph(
                f'匹配段 #{seg_id + 1} — 相似度 {sim_pct}%{cls_label}',
                sub_heading_style,
            ))

            if sc and sc.get('reason'):
                elems.append(Paragraph(
                    f'AI 判定理由：{escape_pdf_text(sc["reason"])}',
                    reason_style,
                ))

            seg_data = [
                [Paragraph(r.get('label_a', '文档 A'), cell_style),
                 Paragraph(r.get('label_b', '文档 B'), cell_style)],
                [Paragraph(escape_pdf_text(seg.get('tender_text', '')), cell_style),
                 Paragraph(escape_pdf_text(seg.get('bid_text', '')), cell_style)],
            ]
            seg_table = Table(seg_data, colWidths=[90 * mm, 90 * mm])
            seg_table.setStyle(TableStyle([
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#888888')),
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e8e8e8')),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('TOPPADDING', (0, 0), (-1, -1), 4),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ]))
            elems.append(seg_table)
            elems.append(Spacer(1, 4 * mm))

        elems.append(PageBreak())

    doc.build(elems)
    buf.seek(0)
    return buf.getvalue()


def escape_pdf_text(text: str) -> str:
    """Escape special characters for reportlab Paragraph."""
    if not text:
        return ''
    return text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
