# -*- coding: utf-8 -*-
"""techtoexcel 通用工具函数"""
import re
import os
from docx import Document


def convert_doc_to_docx(doc_path):
    """将 .doc 文件转换为 .docx 格式（需要安装 pywin32 和 Microsoft Word）

    Args:
        doc_path: .doc 文件的绝对路径

    Returns:
        转换后的 .docx 文件路径，如转换失败返回 None
    """
    try:
        import win32com.client

        # 生成输出路径
        docx_path = doc_path.replace('.doc', '.docx')
        if docx_path == doc_path:
            # 如果扩展名不是.doc结尾，添加.docx
            docx_path = doc_path + '.docx'

        # 转换
        word = win32com.client.Dispatch('Word.Application')
        word.Visible = False

        doc = word.Documents.Open(os.path.abspath(doc_path))
        doc.SaveAs(os.path.abspath(docx_path), FileFormat=16)  # 16 = docx format
        doc.Close()
        word.Quit()

        return docx_path
    except ImportError:
        print("警告: pywin32 未安装，无法转换 .doc 文件")
        print("请执行: pip install pywin32")
        return None
    except Exception as e:
        print(f"警告: Word 转换失败: {e}")
        return None


def ensure_docx(file_path):
    """确保文件为 .docx 格式，如为 .doc 则自动转换

    Args:
        file_path: 输入文件路径

    Returns:
        .docx 格式的文件路径
    """
    if file_path.lower().endswith('.docx'):
        return file_path

    if file_path.lower().endswith('.doc'):
        print(f"检测到 .doc 格式文件，正在转换为 .docx...")
        docx_path = convert_doc_to_docx(file_path)
        if docx_path:
            print(f"转换成功: {docx_path}")
            return docx_path
        else:
            raise ValueError("无法转换 .doc 文件，请手动转换为 .docx 格式或安装 pywin32")

    raise ValueError(f"不支持的文件格式: {file_path}，仅支持 .docx 和 .doc 格式")


def extract_paragraphs(path):
    """提取 docx 全部非空段落，返回 [(索引, 样式名, 文本), ...]"""
    doc = Document(path)
    result = []
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if text:
            result.append((i, p.style.name, text))
    return result


def extract_tables(path):
    """提取 docx 全部表格，返回 [[[单元格文本, ...], ...], ...]"""
    doc = Document(path)
    result = []
    for table in doc.tables:
        rows = []
        for row in table.rows:
            cells = [cell.text for cell in row.cells]
            rows.append(cells)
        result.append(rows)
    return result


def clean_resp(text):
    """清理C列措辞：去★、投标人→我司、去至少/不少于/包括但不限于等"""
    text = text.replace('★', '')
    text = text.replace('投标人', '我司')
    text = text.replace('包括但不限于', '包括')
    text = text.replace('包含且不限于', '包含')
    text = re.sub(r'不少于(\d+)年', r'\1年', text)
    text = re.sub(r'应有\s*(\d+)\s*年', r'\1年', text)
    text = re.sub(r'需(\d+)年', r'\1年', text)
    text = re.sub(r'至少\s*(\d+)次', r'\1次', text)
    text = re.sub(r'至少\s*(\d+)个', r'\1个', text)
    text = text.replace('至少应对', '应对')
    text = text.replace('至少应包含', '应包含')
    text = text.replace('至少应包括', '应包括')
    text = text.replace('至少包括', '包括')
    text = text.replace('至少提供', '提供')
    text = text.replace('要求具有', '具有')
    text = text.replace('要求具备', '具备')
    text = text.replace('至少', '')
    text = re.sub(r'  +', ' ', text)
    text = text.replace('。。', '。').replace('，，', '，')
    return text


def write_excel(entries, output_path):
    """将条目列表写入 Excel"""
    import openpyxl
    from openpyxl.styles import Alignment, Font, Border, Side, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '技术偏离表'

    thin = Border(left=Side('thin'), right=Side('thin'), top=Side('thin'), bottom=Side('thin'))
    hfont = Font(name='宋体', size=11, bold=True)
    cfont = Font(name='宋体', size=10)
    hfill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')

    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 70
    ws.column_dimensions['C'].width = 70
    ws.column_dimensions['D'].width = 15

    for col, h in enumerate(['序号', '招标文件技术条款', '投标文件技术条款', '偏离说明'], 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = hfont; c.fill = hfill; c.border = thin
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for i, e in enumerate(entries):
        r = i + 2
        ws.cell(row=r, column=1, value=i + 1).font = cfont
        ws.cell(row=r, column=1).border = thin
        ws.cell(row=r, column=1).alignment = Alignment(horizontal='center', vertical='top')

        ws.cell(row=r, column=2, value=e['bid']).font = cfont
        ws.cell(row=r, column=2).border = thin
        ws.cell(row=r, column=2).alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)

        ws.cell(row=r, column=3, value=e['resp']).font = cfont
        ws.cell(row=r, column=3).border = thin
        ws.cell(row=r, column=3).alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)

        ws.cell(row=r, column=4, value=e['deviation']).font = cfont
        ws.cell(row=r, column=4).border = thin
        ws.cell(row=r, column=4).alignment = Alignment(horizontal='center', vertical='top')

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:D{len(entries) + 1}'
    wb.save(output_path)
    return len(entries)


def verify(entries):
    """验证条目，返回问题列表"""
    issues = []
    for i, e in enumerate(entries):
        c = e.get('resp', '')
        b = e.get('bid', '')
        if '★' in c:
            issues.append(f'条目{i+1}: C列残留★')
        if '至少' in c:
            issues.append(f'条目{i+1}: C列残留"至少"')
        if '不少于' in c:
            issues.append(f'条目{i+1}: C列残留"不少于"')
        if '包括但不限于' in c:
            issues.append(f'条目{i+1}: C列残留"包括但不限于"')
        if '要求具有' in c:
            issues.append(f'条目{i+1}: C列残留"要求具有"')
        if '投标人' in c:
            issues.append(f'条目{i+1}: C列残留"投标人"')
    return issues
