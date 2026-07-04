from __future__ import annotations

"""File generation service — Markdown and artifact management."""

import os
import re
from datetime import datetime, timezone
from pathlib import Path

from app.config import settings


def sanitize_filename(name: str) -> str:
    """Replace characters unsafe for filenames."""
    unsafe = r'[<>:"/\\|?*\n\r\t]'
    name = re.sub(unsafe, "_", name)
    name = re.sub(r"\s+", " ", name).strip()
    # Limit length
    if len(name) > 200:
        name = name[:200]
    return name


def generate_markdown_summary(
    project_name: str,
    bidder_name: str,
    deadline,
    source_filename: str,
) -> str:
    """Generate a summary markdown file for the bid project.

    Returns the absolute file path of the generated markdown.
    """
    name_safe = sanitize_filename(project_name) if project_name else "未命名项目"
    filename = f"{name_safe}信息V1.md"
    filepath = os.path.join(settings.export_dir, filename)

    deadline_str = ""
    if deadline:
        if hasattr(deadline, "strftime"):
            deadline_str = deadline.strftime("%Y-%m-%d %H:%M")
        else:
            deadline_str = str(deadline)

    now = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")

    content = f"""# {project_name or '未命名项目'} 信息摘要

## 基本信息

| 字段 | 内容 |
|---|---|
| **项目名称** | {project_name or '--'} |
| **招标人名称** | {bidder_name or '--'} |
| **投标截止日期** | {deadline_str or '--'} |
| **源文件** | {source_filename or '--'} |
| **提取时间** | {now} |

---

*本文件由 Bid Info Extractor 自动生成*
"""

    Path(os.path.dirname(filepath)).mkdir(parents=True, exist_ok=True)
    with open(filepath, "w", encoding="utf-8") as f:
        f.write(content)

    return filepath


def generate_tech_excel(
    tech_text: str,
    project_name: str,
) -> str:
    """Generate a basic Excel template for point-by-point technical response.

    This is used when no external techtoexcel script is configured.
    Returns the file path of the generated Excel.
    """
    import openpyxl

    name_safe = sanitize_filename(project_name) if project_name else "未命名项目"
    filename = f"{name_safe}_技术应答.xlsx"
    filepath = os.path.join(settings.export_dir, filename)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "技术点对点应答"

    # Header
    headers = ["序号", "技术要求原文", "应答内容", "偏离说明", "证明材料"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = openpyxl.styles.Font(bold=True)
        cell.fill = openpyxl.styles.PatternFill(
            start_color="4472C4", end_color="4472C4", fill_type="solid"
        )
        cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")

    # Parse tech_text into individual requirements
    # Split by numbered items or paragraph breaks
    items = _parse_requirement_items(tech_text)

    for i, item in enumerate(items):
        row = i + 2
        ws.cell(row=row, column=1, value=i + 1)
        ws.cell(row=row, column=2, value=item)
        ws.cell(row=row, column=3, value="")  # To be filled
        ws.cell(row=row, column=4, value="无偏离")
        ws.cell(row=row, column=5, value="")

    # Column widths
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 50
    ws.column_dimensions["C"].width = 50
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 20

    Path(os.path.dirname(filepath)).mkdir(parents=True, exist_ok=True)
    wb.save(filepath)
    return filepath


def _parse_requirement_items(text: str) -> list[str]:
    """Parse requirement text into individual items."""
    items = []

    # Try splitting by numbered patterns
    pattern = re.compile(
        r"(?:^|\n)\s*(?:\d+[\.\)、]\s*|\(\d+\)\s*|[（(]\d+[）)]\s*)", re.MULTILINE
    )
    splits = pattern.split(text)

    for s in splits:
        s = s.strip()
        if s and len(s) > 10:  # Filter out empty/too-short items
            items.append(s)

    if not items:
        # Fallback: split by paragraphs
        items = [p.strip() for p in text.split("\n") if len(p.strip()) > 10]

    return items[:200]  # Limit to 200 items
