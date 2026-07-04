"""Document parsing service — PDF/DOCX with section structure extraction."""

from __future__ import annotations

import os
import re
from dataclasses import dataclass, field
from typing import Optional


@dataclass
class Section:
    """A document section with title and content."""

    title: str
    level: int = 1  # 1=Heading1, 2=Heading2...
    content: str = ""
    start_index: int = 0
    end_index: int = 0


@dataclass
class ParsedDocument:
    """Parsed document with paragraphs and section structure."""

    filename: str
    full_text: str
    paragraphs: list[dict] = field(default_factory=list)
    sections: list[Section] = field(default_factory=list)
    metadata: dict = field(default_factory=dict)


def parse_file(filepath: str, original_filename: Optional[str] = None) -> ParsedDocument:
    """Parse a file, dispatching by extension."""
    ext = os.path.splitext(filepath)[1].lower()
    filename = original_filename or os.path.basename(filepath)

    if ext == ".docx":
        return _parse_docx(filepath, filename)
    elif ext == ".pdf":
        return _parse_pdf(filepath, filename)
    else:
        raise ValueError(f"不支持的文件格式: {ext}")


# ── Heading detection helpers ──────────────────────────────

_HEADING_PATTERNS = [
    re.compile(r"^第[一二三四五六七八九十百千\d]+章\s*[\.、\s]?\s*.*"),
    re.compile(r"^第[一二三四五六七八九十百千\d]+节\s*[\.、\s]?\s*.*"),
    re.compile(r"^[一二三四五六七八九十]+[、．.\s].*"),
    re.compile(r"^\d+[\.、\s]+[^\d].*"),
    re.compile(r"^（[一二三四五六七八九十\d]+）.*"),
]


def _is_heading(text: str, style: str = "") -> bool:
    """Detect if a paragraph is a heading based on style or content patterns."""
    if style and ("Heading" in style or "heading" in style or "Title" in style):
        return True
    text = text.strip()
    if len(text) > 50:
        return False
    for pat in _HEADING_PATTERNS:
        if pat.match(text):
            return True
    return False


def _get_heading_level(text: str, style: str = "") -> int:
    """Infer heading level from style name or content."""
    if style:
        for level in range(1, 7):
            if f"Heading{level}" in style or f"heading{level}" in style:
                return level
    if re.match(r"^第[一二三四五六七八九十百千\d]+章", text):
        return 1
    if re.match(r"^第[一二三四五六七八九十百千\d]+节", text):
        return 2
    return 1


def _build_sections(paragraphs: list[dict]) -> list[Section]:
    """Group paragraphs into sections by detected headings."""
    sections: list[Section] = []
    current_section: Section | None = None
    current_content: list[str] = []

    for i, p in enumerate(paragraphs):
        is_heading = p.get("is_heading", False)
        if is_heading:
            # Save previous section
            if current_section is not None:
                current_section.content = "\n".join(current_content).strip()
                current_section.end_index = i - 1
                sections.append(current_section)
                current_content = []

            current_section = Section(
                title=p["text"],
                level=_get_heading_level(p["text"], p.get("style", "")),
                start_index=i,
            )
        elif current_section is not None:
            current_content.append(p["text"])
        else:
            # Content before any heading → preamble section
            current_section = Section(title="(文档开头)", level=0, start_index=0)
            current_content.append(p["text"])

    # Save final section
    if current_section is not None:
        current_section.content = "\n".join(current_content).strip()
        current_section.end_index = len(paragraphs) - 1
        sections.append(current_section)

    return sections


# ── Format-specific parsers ─────────────────────────────────


def _parse_docx(filepath: str, filename: str) -> ParsedDocument:
    """Parse .docx with heading-aware structure extraction."""
    from docx import Document

    doc = Document(filepath)
    paragraphs: list[dict] = []
    idx = 0

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = para.style.name if para.style else "Normal"
        is_heading = _is_heading(text, style)
        paragraphs.append({
            "text": text,
            "style": style,
            "page": None,
            "index": idx,
            "is_heading": is_heading,
        })
        idx += 1

    # Tables
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                row_text = " | ".join(cells)
                paragraphs.append({
                    "text": row_text,
                    "style": "table",
                    "page": None,
                    "index": idx,
                    "is_heading": False,
                })
                idx += 1

    full_text = "\n".join(p["text"] for p in paragraphs)
    sections = _build_sections(paragraphs)

    return ParsedDocument(
        filename=filename,
        full_text=full_text,
        paragraphs=paragraphs,
        sections=sections,
        metadata={
            "format": "docx",
            "chars": len(full_text),
            "paragraphs": len(paragraphs),
            "sections": len(sections),
            "tables": len(doc.tables),
        },
    )


def _parse_pdf(filepath: str, filename: str) -> ParsedDocument:
    """Parse PDF with heading detection."""
    import pdfplumber

    paragraphs: list[dict] = []
    idx = 0

    with pdfplumber.open(filepath) as pdf:
        for page_num, page in enumerate(pdf.pages, 1):
            text = page.extract_text()
            if text:
                for line in text.strip().split("\n"):
                    line = line.strip()
                    if not line:
                        continue
                    is_heading = _is_heading(line)
                    paragraphs.append({
                        "text": line,
                        "style": "Heading" if is_heading else "Normal",
                        "page": page_num,
                        "index": idx,
                        "is_heading": is_heading,
                    })
                    idx += 1

            # Tables
            for table in page.extract_tables():
                for row in table:
                    cells = [c.strip() if c else "" for c in row]
                    row_text = " | ".join(c for c in cells if c)
                    if row_text:
                        paragraphs.append({
                            "text": row_text,
                            "style": "table",
                            "page": page_num,
                            "index": idx,
                            "is_heading": False,
                        })
                        idx += 1

    full_text = "\n".join(p["text"] for p in paragraphs)
    sections = _build_sections(paragraphs)

    return ParsedDocument(
        filename=filename,
        full_text=full_text,
        paragraphs=paragraphs,
        sections=sections,
        metadata={
            "format": "pdf",
            "chars": len(full_text),
            "paragraphs": len(paragraphs),
            "sections": len(sections),
            "pages": len(pdf.pages),
        },
    )
